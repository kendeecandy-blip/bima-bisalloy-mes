import os
import time
import math
import io
import re
import pytz
from datetime import datetime, timedelta, date
import streamlit as st
import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.exc import IntegrityError
from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

# --- IMPOR MODUL LOKAL ---
from models import (
    Base, Customer, Vendor, ScopeMaster, Project, 
    ProjectScopeRouting, ProjectProgressLog, QCStatus, 
    VendorType, ProjectStatus, get_jakarta_time
)

# --- SETTING TEMA & LOGO KORPORAT ---
st.set_page_config(page_title="Bima Bisalloy - Smart MES & Tonnage Tracker", page_icon="🏗️", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
    <style>
    /* Global & Button */
    .stButton>button { background-color: #003366; color: white; border: none; font-weight: bold; }
    .stButton>button:hover { background-color: #d92323; color: white; }
    
    /* Dashboard Metric Cards */
    .metric-card { background-color: #f8f9fa; border-left: 5px solid #003366; padding: 15px; border-radius: 8px; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.05); }
    .metric-card h4 { margin: 0; font-size: 15px; color: #666; font-weight: normal; }
    .metric-card h2 { margin: 5px 0 0 0; font-size: 32px; color: #003366; font-weight: bold; }
    
    /* SIDEBAR THEME FIX - Dark Mode & White Text */
    [data-testid="stSidebar"] { background-color: #002244 !important; }
    [data-testid="stSidebar"] * { color: #FFFFFF !important; }
    /* Khusus mengatasi radio button/checkbox text di sidebar */
    [data-testid="stSidebar"] .stRadio label p, [data-testid="stSidebar"] .stMarkdown p { color: #FFFFFF !important; }
    </style>
""", unsafe_allow_html=True)

# --- DATABASE & KONFIGURASI DIREKTORI ---
load_dotenv()
DB_URL = os.getenv("DATABASE_URL", "sqlite:///bima_bisalloy_qc.db")
engine = create_engine(DB_URL, echo=False)
SessionLocal = sessionmaker(bind=engine)
Base.metadata.create_all(bind=engine)

UPLOAD_DIR = "uploaded_images"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# --- HELPER PARSING TONASE ---
def parse_tonnage_from_desc(desc):
    desc_str = str(desc)
    match = re.search(r'(\d+(?:\.\d+)?)\s*(PCS|UNIT|SET|EA|KG|TON)', desc_str, re.IGNORECASE)
    if match:
        val = float(match.group(1))
        unit = match.group(2).upper()
        if unit == 'KG': return val
        elif unit == 'TON': return val * 1000.0
        else: return val * 150.0  
    return 1000.0

# --- MESIN PARSING EXCEL & INISIALISASI DATA ---
def init_db_data():
    db = SessionLocal()
    v_int = db.query(Vendor).filter(Vendor.name == "Workshop Fabrikasi Internal").first()
    if not v_int:
        try:
            v_int = Vendor(name="Workshop Fabrikasi Internal", type=VendorType.INTERNAL)
            db.add(v_int)
            db.commit()
            db.refresh(v_int)
        except IntegrityError:
            db.rollback()
            v_int = db.query(Vendor).filter(Vendor.name == "Workshop Fabrikasi Internal").first()
        
    excel_file_path = 'data_manual_bima.xlsx'
    if os.path.exists(excel_file_path):
        try:
            df = pd.read_excel(excel_file_path, skiprows=3)
            for idx, row in df.iterrows():
                if pd.isna(row.get('CUSTOMER')) or pd.isna(row.get('SO NO.')) or pd.isna(row.get('DESCRIPTION')): continue
                
                raw_no = str(row.get('NO.', '')).strip().upper()
                raw_cust = str(row.get('CUSTOMER', '')).strip()
                raw_so = str(row.get('SO NO.', '')).strip()
                raw_desc = str(row.get('DESCRIPTION', '')).strip()

                if not raw_cust or raw_cust.lower() == 'nan': continue
                if not raw_so or raw_so.lower() == 'nan': continue
                if not raw_desc or raw_desc.lower() == 'nan': continue

                recap_keywords = ['TOTAL', 'OUTSTANDING', 'O', 'STEEL', 'FERRITIC']
                skip_row = False
                for kw in recap_keywords:
                    if kw == 'O':
                        if kw == raw_no or f" {kw} " in f" {raw_no} ": skip_row = True
                    else:
                        if kw in raw_no: skip_row = True
                
                if skip_row: continue

                if raw_cust.startswith("('") and raw_cust.endswith("',)"): raw_cust = eval(raw_cust)[0]

                cust = db.query(Customer).filter(Customer.name == raw_cust).first()
                if not cust:
                    try:
                        cust = Customer(name=raw_cust)
                        db.add(cust)
                        db.commit()
                        db.refresh(cust)
                    except IntegrityError:
                        db.rollback()
                        cust = db.query(Customer).filter(Customer.name == raw_cust).first()

                if raw_so.endswith(".0"): raw_so = raw_so.replace(".0", "")
                if raw_so.startswith("('"): raw_so = eval(raw_so)[0]
                if db.query(Project).filter(Project.so_number == raw_so).first(): continue 
                
                yej_cycle_val = 'YEJ26' if raw_so.startswith('YEJ26') else 'YEJ27'
                raw_mat = str(row.get('RAW MATERIAL', '')).upper()
                mat_list = []
                if 'WEAR PLATE' in raw_mat: mat_list.append("FAB01 - Wear Plate")
                if 'STRUCTURAL' in raw_mat: mat_list.append("FAB03 - Structural")
                if 'SS400' in raw_mat: mat_list.append("FAB06 - Carbon Steel")
                if 'STAINLESS STEEL' in raw_mat: mat_list.append("FAB04 - Stainless Steel")
                final_material = ",".join(mat_list) if mat_list else raw_mat

                if raw_desc.startswith("('") and raw_desc.endswith("',)"): raw_desc = eval(raw_desc)[0]
                final_proj_name = raw_desc if raw_desc else f"Project {raw_so}"
                migrated_tonnage = parse_tonnage_from_desc(raw_desc)

                raw_status = re.sub(r'[^A-Za-z0-9\s]', '', str(row.get('STATUS', ''))).strip().upper()
                proj_status_val = ProjectStatus.COMPLETED if raw_status == 'CLOSED' else ProjectStatus.IN_PRODUCTION

                logistic_status_val = None
                delivery_num_val = str(row.get('Delivery Number', '')).strip()
                if delivery_num_val.lower() == 'nan': delivery_num_val = ''
                goods_arrived_val = str(row.get('GOODS ARRIVED', '')).strip()
                if goods_arrived_val.lower() == 'nan': goods_arrived_val = ''

                if raw_status == 'CLOSED':
                    if not delivery_num_val: logistic_status_val = "🏭 FABRIKASI SELESAI (BELUM DIKIRIM)"
                    elif not goods_arrived_val: logistic_status_val = "🚚 DALAM PERJALANAN (BELUM DITERIMA)"
                    else: logistic_status_val = "✅ SELESAI TOTAL (SUDAH DITERIMA KONSUMEN)"
                elif 'ON PROGRESS' in raw_status:
                    logistic_status_val = "🏗️ PROSES FABRIKASI"

                now_local = get_jakarta_time()
                new_proj = Project(
                    customer_id=cust.id, project_name=final_proj_name, so_number=raw_so, sales_pic="Sistem Migrasi",
                    total_tonnage=migrated_tonnage, material_list=final_material, is_jasa_only=False, yej_cycle=yej_cycle_val,
                    start_date=(now_local - timedelta(days=30)).date(), target_selesai_global=(now_local + timedelta(days=15)).date(),
                    status_proyek=proj_status_val, logistic_status=logistic_status_val, delivery_number=delivery_num_val if delivery_num_val else None
                )
                db.add(new_proj)
                db.commit()
                db.refresh(new_proj)

                raw_proc = str(row.get('PROCESS', ''))
                proc_items = raw_proc.split(';')
                scopes_in_order = []
                for p in proc_items:
                    if not p.strip(): continue
                    vendor_match = re.search(r'\(MATERIAL DARI VENDOR:\s*(.*?)\)', p, re.IGNORECASE)
                    if vendor_match:
                        ext_v_name = vendor_match.group(1).strip()
                        v_ext = db.query(Vendor).filter(Vendor.name == ext_v_name).first()
                        if not v_ext:
                            try:
                                v_ext = Vendor(name=ext_v_name, type=VendorType.EXTERNAL)
                                db.add(v_ext)
                                db.commit()
                                db.refresh(v_ext)
                            except IntegrityError:
                                db.rollback()
                                v_ext = db.query(Vendor).filter(Vendor.name == ext_v_name).first()
                    
                    clean_scope_name = re.sub(r'\(.*?\)', '', p).strip().upper()
                    if clean_scope_name:
                        sc = db.query(ScopeMaster).filter(ScopeMaster.name == clean_scope_name).first()
                        if not sc:
                            try:
                                sc = ScopeMaster(name=clean_scope_name)
                                db.add(sc)
                                db.commit()
                                db.refresh(sc)
                            except IntegrityError:
                                db.rollback()
                                sc = db.query(ScopeMaster).filter(ScopeMaster.name == clean_scope_name).first()
                        
                        db.add(ProjectScopeRouting(project_id=new_proj.id, scope_id=sc.id))
                        scopes_in_order.append(sc)
                db.commit()

                if raw_status == 'CLOSED':
                    for sc in scopes_in_order:
                        db.add(ProjectProgressLog(project_id=new_proj.id, scope_id=sc.id, vendor_id=v_int.id, actual_tonnage_completed=migrated_tonnage, status=QCStatus.PASS, qc_notes="Auto-Migration Closed", log_date=now_local - timedelta(days=5)))
                elif 'ON PROGRESS' in raw_status:
                    clean_status = raw_status.replace('ON PROGRESS', '').strip()
                    pct_match = re.search(r'([A-Za-z\s]+?)\s*(\d+)\s*%', clean_status)
                    if pct_match:
                        curr_scope = pct_match.group(1).strip().upper()
                        pct = float(pct_match.group(2))
                        found_active = False
                        log_date_sim = now_local - timedelta(days=10)
                        
                        for sc in scopes_in_order:
                            if sc.name == curr_scope:
                                found_active = True
                                ton_done = (pct / 100.0) * migrated_tonnage
                                db.add(ProjectProgressLog(project_id=new_proj.id, scope_id=sc.id, vendor_id=v_int.id, actual_tonnage_completed=ton_done, status=QCStatus.PASS, qc_notes=f"Auto-Migration In-Progress ({pct}%)", log_date=log_date_sim))
                            elif not found_active:
                                db.add(ProjectProgressLog(project_id=new_proj.id, scope_id=sc.id, vendor_id=v_int.id, actual_tonnage_completed=migrated_tonnage, status=QCStatus.PASS, qc_notes="Auto-Migration Done", log_date=log_date_sim))
                            else:
                                db.add(ProjectProgressLog(project_id=new_proj.id, scope_id=sc.id, vendor_id=v_int.id, actual_tonnage_completed=0.0, status=QCStatus.PASS, qc_notes="Auto-Migration Pending", log_date=log_date_sim))
                            log_date_sim += timedelta(days=1)
                db.commit()
        except Exception as e:
            print(f"Error parsing Excel: {e}")
            db.rollback()
    db.close()

init_db_data()

# --- SISTEM LOGIN & MATRIKS HAK AKSES ---
def display_logo():
    if os.path.exists("logo_bima.png"): st.image("logo_bima.png", width=250)
    else: st.markdown("<h2 style='color: #FFFFFF;'>🏗️ PT BIMA BISALLOY</h2>", unsafe_allow_html=True)

def login_screen():
    col1, col2, col3 = st.columns([1,2,1])
    with col2:
        if os.path.exists("logo_bima.png"): st.image("logo_bima.png", width=250)
        st.subheader("🔐 Enterprise MES & Tonnage Tracker")
        with st.form("login_form"):
            username = st.text_input("Username")
            password = st.text_input("Password", type="password")
            submitted = st.form_submit_button("Login to System")
            
            if submitted:
                if username == "ppic" and password == "111":
                    st.session_state["logged_in"] = True; st.session_state["role"] = "PPIC (Super User)"
                    st.session_state["access"] = ["Registrasi Proyek", "Perencanaan Target", "Gambar Kerja", "Update Progress & QC", "Update Status Logistik", "Management Dashboard"]
                    st.rerun()
                elif username == "qc" and password == "2623":
                    st.session_state["logged_in"] = True; st.session_state["role"] = "QC Lapangan"
                    st.session_state["access"] = ["Update Progress & QC"]
                    st.rerun()
                elif username == "management" and password == "6789":
                    st.session_state["logged_in"] = True; st.session_state["role"] = "Management"
                    st.session_state["access"] = ["Update Status Logistik", "Management Dashboard"]
                    st.rerun()
                else:
                    st.error("Kredensial tidak valid sesuai Matriks Hak Akses.")

def logout():
    st.session_state.clear()
    st.rerun()

# --- MODULE 1: ADMIN SO / PPIC (Registrasi & Dynamic Routing) ---
def view_registrasi_proyek():
    st.header("📝 Registrasi Proyek & SO Baru")
    db = SessionLocal()

    with st.expander("➕ Tambah Jenis Proses Fabrikasi Baru ke Master Data"):
        new_scope_input = st.text_input("Nama Proses Fabrikasi Baru")
        if st.button("Simpan ke Master Scope") and new_scope_input.strip():
            clean_name = new_scope_input.strip().upper()
            if not db.query(ScopeMaster).filter(ScopeMaster.name == clean_name).first():
                db.add(ScopeMaster(name=clean_name))
                db.commit()
                st.success(f"Proses '{clean_name}' berhasil ditambahkan ke Master Data!")
                st.rerun()
            else: st.warning("Proses tersebut sudah ada di Master Data.")

    st.markdown("---")
    with st.form("register_project_form"):
        customer_options = ["[+ Ketik Customer Baru...]"] + [c.name for c in db.query(Customer).all()]
        
        st.markdown("**1. Informasi Pelanggan & SO**")
        colA, colB = st.columns(2)
        selected_cust_name = colA.selectbox("Pilih Customer", customer_options)
        new_cust_name = colB.text_input("📝 Masukkan Nama Customer Baru") if selected_cust_name == "[+ Ketik Customer Baru...]" else ""
            
        raw_so_num = colA.text_input("SO Number / Nama Tiket SO").strip()
        pic = colB.text_input("Sales PIC")
        
        project_options = ["[+ Ketik Proyek Baru...]"] + [p.project_name for p in db.query(Project.project_name).distinct().all() if p.project_name]
        
        st.markdown("**2. Detail Proyek, Material, & Tonase**")
        colC, colD = st.columns(2)
        selected_proj_name = colC.selectbox("Pilih Nama Proyek", project_options)
        new_proj_name = colD.text_input("📝 Masukkan Nama Proyek Baru") if selected_proj_name == "[+ Ketik Proyek Baru...]" else ""
        total_tonnage = colC.number_input("Total Berat Material Proyek (Kg/Ton)", min_value=0.1, step=0.1)
        
        material_opts = ["FAB01 - Wear Plate", "FAB02 - Cladding", "FAB03 - Structural", "FAB04 - Stainless Steel", "FAB05 - Roundbar", "FAB06 - Carbon Steel"]
        selected_materials = colD.multiselect("Jenis Material Fabrikasi", options=material_opts)
        is_jasa_only_val = st.checkbox("Proyek Ini Hanya Jasa (Tanpa Pengadaan Material)?")

        st.markdown("**3. Routing Tahapan Proses Fabrikasi Aktif**")
        selected_scopes = st.multiselect("Pilih Tahapan Fabrikasi", options=[s.name for s in db.query(ScopeMaster).all()])
        
        if st.form_submit_button("Daftarkan Proyek & Lock Routing"):
            if raw_so_num.isdigit() and len(raw_so_num) == 8: so_num = f"#{raw_so_num}"
            else: so_num = raw_so_num

            final_proj_name = new_proj_name.strip() if selected_proj_name == "[+ Ketik Proyek Baru...]" else selected_proj_name
            if not final_proj_name: st.error("❌ Nama Proyek wajib diisi!")
            elif not so_num: st.error("❌ Nomor SO wajib diisi!")
            else:
                final_cust_id = None
                if selected_cust_name == "[+ Ketik Customer Baru...]" and new_cust_name.strip():
                    try:
                        new_cust = Customer(name=new_cust_name.strip())
                        db.add(new_cust)
                        db.commit()
                        final_cust_id = new_cust.id
                    except IntegrityError:
                        db.rollback()
                        final_cust_id = db.query(Customer).filter(Customer.name == new_cust_name.strip()).first().id
                elif selected_cust_name != "[+ Ketik Customer Baru...]":
                    final_cust_id = db.query(Customer).filter(Customer.name == selected_cust_name).first().id
                
                yej_cycle_val = 'YEJ26' if so_num.startswith('YEJ26') else 'YEJ27'

                if final_cust_id:
                    new_proj = Project(
                        customer_id=final_cust_id, project_name=final_proj_name, so_number=so_num, sales_pic=pic, total_tonnage=float(total_tonnage), 
                        material_list=",".join(selected_materials), is_jasa_only=is_jasa_only_val, yej_cycle=yej_cycle_val, status_proyek=ProjectStatus.WAITING_SCHEDULE
                    )
                    try:
                        db.add(new_proj)
                        db.commit()
                        db.refresh(new_proj)
                        final_scopes_set = set(selected_scopes) | {"FINISHING", "PACKING", "DELIVERY", "RECEIVED"}
                        for scope_name in final_scopes_set:
                            sc_obj = db.query(ScopeMaster).filter(ScopeMaster.name == scope_name).first()
                            if not sc_obj:
                                sc_obj = ScopeMaster(name=scope_name)
                                db.add(sc_obj)
                                db.commit()
                            db.add(ProjectScopeRouting(project_id=new_proj.id, scope_id=sc_obj.id))
                        db.commit()
                        st.success(f"Berhasil! Proyek '{final_proj_name}' ({so_num}) terdaftar.")
                    except IntegrityError:
                        db.rollback(); st.error(f"Gagal! SO Number {so_num} sudah terdaftar.")
    db.close()

# --- MODULE 2: PPIC ONLY (Perencanaan Target) ---
def view_perencanaan_target():
    st.header("📅 Perencanaan Jadwal & Target Global (PPIC)")
    db = SessionLocal()
    projects = db.query(Project).filter(Project.status_proyek == ProjectStatus.WAITING_SCHEDULE).all()
    if not projects: st.info("Tidak ada proyek yang menunggu penjadwalan saat ini.")
    else:
        with st.form("target_form"):
            selected_proj = st.selectbox("Pilih Proyek", projects, format_func=lambda x: f"{x.so_number} - {x.project_name} ({x.total_tonnage} Kg/Ton)")
            col1, col2 = st.columns(2)
            start_date = col1.date_input("Mulai Produksi (Baseline)")
            target_date = col2.date_input("Target Selesai Produksi Global")
            
            if st.form_submit_button("Set Target & Majukan ke 'Waiting Drawing'"):
                if target_date < start_date: st.error("❌ Target selesai tidak boleh mendahului tanggal mulai!")
                else:
                    selected_proj.start_date = start_date
                    selected_proj.target_selesai_global = target_date
                    selected_proj.status_proyek = ProjectStatus.WAITING_DRAWING
                    db.commit()
                    st.success(f"Rentang jadwal SO {selected_proj.so_number} ditetapkan. Dilempar ke Engineering.")
    db.close()

# --- MODULE 3: ENGINEERING / PPIC (Gambar Kerja) ---
def view_gambar_kerja():
    st.header("📐 Approval Gambar Kerja & Desain")
    db = SessionLocal()
    projects = db.query(Project).filter(Project.status_proyek == ProjectStatus.WAITING_DRAWING).all()
    if not projects: st.info("Semua proyek sudah masuk tahap produksi atau belum dijadwalkan PPIC.")
    else:
        with st.form("drawing_form"):
            selected_proj = st.selectbox("Pilih Proyek", projects, format_func=lambda x: f"{x.so_number} - {x.project_name}")
            approved = st.checkbox("Gambar Kerja Approved untuk Produksi?")
            notes = st.text_area("Catatan Desain / Dimensi Khusus")
            
            if st.form_submit_button("Rilis ke Produksi"):
                if approved:
                    selected_proj.drawing_approved = True
                    selected_proj.drawing_notes = notes
                    selected_proj.status_proyek = ProjectStatus.IN_PRODUCTION
                    db.commit()
                    st.success(f"Desain SO {selected_proj.so_number} dirilis! Proyek sekarang 'In Production'.")
                else: st.warning("Centang kotak 'Approved' untuk merilis ke produksi.")
    db.close()

# --- MODULE 4: QC / PPIC (Update Progress Fisik & Arsip Foto) ---
def view_update_qc():
    st.header("⚡ Update Progress & Manajemen Log Lapangan")
    tab1, tab2 = st.tabs(["Update Progress Tonnage (Live)", "Arsip Foto Susulan Data Migrasi"])
    db = SessionLocal()
    
    # ---------------- TAB 1: UPDATE LIVE PROGRESS ----------------
    with tab1:
        projects = db.query(Project).filter(Project.status_proyek == ProjectStatus.IN_PRODUCTION).all()
        if not projects:
            st.warning("Tidak ada proyek di tahap produksi saat ini.")
        else:
            selected_project = st.selectbox("Pilih Proyek Active", projects, format_func=lambda x: f"{x.so_number} - {x.customer.name} ({x.project_name})", key="qc_proj_sel")
            routings = db.query(ProjectScopeRouting).filter(ProjectScopeRouting.project_id == selected_project.id).order_by(ProjectScopeRouting.id.asc()).all()
            active_scopes = [r.scope for r in routings]

            if not active_scopes:
                st.error("Proyek ini belum memiliki routing pekerjaan.")
            else:
                default_scope_idx = 0
                if "auto_advance_map" in st.session_state and selected_project.id in st.session_state["auto_advance_map"]:
                    target_scope_id = st.session_state["auto_advance_map"][selected_project.id]
                    for idx, sc in enumerate(active_scopes):
                        if sc.id == target_scope_id: default_scope_idx = idx; break

                selected_scope = st.selectbox("Tahapan Scope (Reaktif Sesuai Routing Proyek)", active_scopes, index=default_scope_idx, format_func=lambda x: x.name)
                
                cutting_tech = None
                if selected_scope.name.upper() == "CUTTING":
                    cutting_tech = st.selectbox("Jenis Teknologi Cutting yang Digunakan (Sub-Scope)", ["Oxy-Fuel Cutting", "Plasma Cutting", "Laser Cutting"])

                db_vendors = db.query(Vendor).all()
                vendor_opts = ["[+ Ketik Vendor Baru...]"] + db_vendors
                selected_vendor = st.selectbox("Vendor Pelaksana", vendor_opts, format_func=lambda x: x if isinstance(x, str) else f"{x.name} ({x.type.value})")
                
                new_vendor_name = ""
                is_external = False
                if selected_vendor == "[+ Ketik Vendor Baru...]":
                    new_vendor_name = st.text_input("📝 Masukkan Nama Vendor Baru (Otomatis diset sebagai EXTERNAL)")
                    is_external = True
                else: is_external = (selected_vendor.type == VendorType.EXTERNAL)
                
                st.markdown("---")
                with st.form("qc_update_form"):
                    st.markdown("**📸 Upload Foto Bukti Lapangan**")
                    uploaded_file = st.file_uploader("Lampirkan Foto Bukti Visual Fisik (PNG, JPG, JPEG)", type=["png", "jpg", "jpeg"])
                    caption_live = st.text_input("Keterangan Foto (Opsional)")
                    po_vendor_val = st.text_input("Nomor PO Vendor (Mandatory untuk Vendor External)") if is_external else None
                    mat_label = "✅ Validasi Material Customer: Dipastikan Sesuai Spesifikasi Jasa" if selected_project.is_jasa_only else "✅ Validasi Material Gudang: Sesuai Spesifikasi?"
                    material_ok = st.checkbox(mat_label)
                    despatch_val = st.text_input("Despatch to BB (No Resi/Surat Jalan)") if is_external else None
                    
                    col1, col2 = st.columns(2)
                    tonase_hari_ini = col1.number_input("Berat Tonase Material yang Selesai Hari Ini (Kg/Ton)", min_value=0.0, step=0.1)
                    status_val = col2.selectbox("Status Hasil Inspeksi", ["Pass", "Fail", "Hold"])
                    notes = st.text_area("Catatan Temuan QC / Kendala")
                    
                    if st.form_submit_button("Submit Data Tonase Fisik"):
                        all_logs_scope = db.query(ProjectProgressLog).filter(ProjectProgressLog.project_id == selected_project.id, ProjectProgressLog.scope_id == selected_scope.id, ProjectProgressLog.status == QCStatus.PASS).all()
                        total_scope_tonnage = sum(l.actual_tonnage_completed for l in all_logs_scope)

                        if status_val == "Pass" and (total_scope_tonnage + tonase_hari_ini) > selected_project.total_tonnage:
                            st.error(f"❌ OVER-TONNAGE ERROR! Akumulasi lulus saat ini ({total_scope_tonnage}) ditambah input hari ini ({tonase_hari_ini}) melampaui Total Target Proyek ({selected_project.total_tonnage}). Proses Dibatalkan.")
                        elif is_external and not po_vendor_val.strip(): st.error("❌ Nomor PO Vendor WAJIB diisi untuk Vendor Eksternal!")
                        elif selected_vendor == "[+ Ketik Vendor Baru...]" and not new_vendor_name.strip(): st.error("❌ Nama Vendor Baru tidak boleh kosong!")
                        elif status_val in ["Fail", "Hold"] and not notes.strip(): st.error("❌ Catatan Temuan WAJIB diisi jika status Fail atau Hold!")
                        else:
                            final_vendor_id = None
                            if selected_vendor == "[+ Ketik Vendor Baru...]":
                                existing_v = db.query(Vendor).filter(Vendor.name == new_vendor_name.strip()).first()
                                if existing_v: final_vendor_id = existing_v.id
                                else:
                                    try:
                                        new_v = Vendor(name=new_vendor_name.strip(), type=VendorType.EXTERNAL)
                                        db.add(new_v); db.commit(); db.refresh(new_v)
                                        final_vendor_id = new_v.id
                                    except IntegrityError:
                                        db.rollback()
                                        final_vendor_id = db.query(Vendor).filter(Vendor.name == new_vendor_name.strip()).first().id
                            else: final_vendor_id = selected_vendor.id

                            saved_image_path = None
                            now_local = get_jakarta_time()
                            if uploaded_file is not None:
                                timestamp = now_local.strftime("%Y%m%d_%H%M%S")
                                clean_so = selected_project.so_number.replace("#", "").replace("/", "_")
                                saved_image_path = os.path.join(UPLOAD_DIR, f"{clean_so}_{timestamp}.{uploaded_file.name.split('.')[-1]}")
                                with open(saved_image_path, "wb") as f: f.write(uploaded_file.getbuffer())

                            final_notes = notes.strip()
                            if cutting_tech: final_notes = f"[Mesin: {cutting_tech}] - {final_notes}" if final_notes else f"[Mesin: {cutting_tech}]"

                            new_log = ProjectProgressLog(
                                project_id=selected_project.id, scope_id=selected_scope.id, vendor_id=final_vendor_id,
                                actual_tonnage_completed=tonase_hari_ini, qc_notes=final_notes, status=QCStatus[status_val.upper()],
                                material_validated=material_ok, despatch_to_bb=despatch_val, image_path=saved_image_path,
                                image_caption=caption_live if caption_live else None, po_vendor_number=po_vendor_val, log_date=now_local
                            )
                            db.add(new_log)
                            db.commit()
                            
                            total_scope_tonnage += (tonase_hari_ini if status_val == "Pass" else 0)
                            if total_scope_tonnage >= selected_project.total_tonnage:
                                curr_idx = -1
                                for idx, sc in enumerate(active_scopes):
                                    if sc.id == selected_scope.id: curr_idx = idx; break
                                if curr_idx != -1:
                                    if curr_idx + 1 < len(active_scopes):
                                        next_scope_obj = active_scopes[curr_idx + 1]
                                        if "auto_advance_map" not in st.session_state: st.session_state["auto_advance_map"] = {}
                                        st.session_state["auto_advance_map"][selected_project.id] = next_scope_obj.id
                                    else:
                                        selected_project.status_proyek = ProjectStatus.COMPLETED
                                        selected_project.actual_delivery_date = now_local.date()
                                        db.commit()

                            st.success("✅ Progress Tonase Fisik & Info Vendor Tersimpan di Log Transaksi.")
                            time.sleep(1.5); st.rerun()

    # ---------------- TAB 2: ARSIP FOTO SUSULAN ----------------
    with tab2:
        st.subheader("📸 Sisipkan Foto Bukti pada Log Migrasi Terdahulu")
        all_projs = db.query(Project).all()
        if not all_projs:
            st.info("Belum ada proyek terdaftar."); db.close(); return
            
        sel_proj = st.selectbox("Pilih Proyek (Arsip Foto)", all_projs, format_func=lambda x: f"{x.so_number} - {x.project_name}")
        proj_logs = db.query(ProjectProgressLog).filter(ProjectProgressLog.project_id == sel_proj.id).order_by(ProjectProgressLog.log_date.asc()).all()
        
        if not proj_logs: st.info("Belum ada log progress untuk proyek ini.")
        else:
            with st.form("arsip_foto_form"):
                log_opts = {l.id: f"{l.log_date.strftime('%Y-%m-%d %H:%M')} | {l.scope.name} | {l.actual_tonnage_completed} Kg | Status: {l.status.value}" for l in proj_logs}
                sel_log_id = st.selectbox("Pilih Log Progress yang Ingin Disisipkan Foto", options=list(log_opts.keys()), format_func=lambda x: log_opts[x])
                
                up_foto = st.file_uploader("Unggah Foto Bukti Baru (PNG, JPG, JPEG)", type=["png", "jpg", "jpeg"])
                caption_text = st.text_input("Keterangan Foto Bukti (Manual Input)")
                
                if st.form_submit_button("Simpan Foto ke Arsip Log"):
                    target_log = db.query(ProjectProgressLog).filter(ProjectProgressLog.id == sel_log_id).first()
                    if target_log:
                        if up_foto is not None:
                            now_local = get_jakarta_time()
                            timestamp = now_local.strftime("%Y%m%d_%H%M%S")
                            clean_so = sel_proj.so_number.replace("#", "").replace("/", "_")
                            saved_path = os.path.join(UPLOAD_DIR, f"{clean_so}_arsip_{timestamp}.{up_foto.name.split('.')[-1]}")
                            with open(saved_path, "wb") as f: f.write(up_foto.getbuffer())
                            target_log.image_path = saved_path
                        if caption_text.strip():
                            target_log.image_caption = caption_text.strip()
                        db.commit()
                        st.success("✅ Foto dan Keterangan berhasil ditambahkan pada log tersebut!")
                        time.sleep(1.5); st.rerun()
                    else: st.error("Log tidak ditemukan.")
    db.close()

# --- MODULE 4.5: UPDATE LOGISTIK & PENGIRIMAN ---
def view_update_logistik():
    st.header("🚚 Update Status Logistik & Pengiriman")
    db = SessionLocal()
    
    projects = db.query(Project).filter(Project.status_proyek == ProjectStatus.COMPLETED).all()
    if not projects:
        st.info("Belum ada proyek yang berstatus COMPLETED.")
        db.close(); return
        
    with st.form("logistic_form"):
        selected_proj = st.selectbox("Pilih Proyek (COMPLETED)", projects, format_func=lambda x: f"{x.so_number} - {x.project_name}")
        status_opts = ["🏭 FABRIKASI SELESAI (BELUM DIKIRIM)", "🚚 DALAM PERJALANAN (BELUM DITERIMA)", "✅ SELESAI TOTAL (SUDAH DITERIMA KONSUMEN)", "Lainnya (Ketik Bebas)"]
        
        col1, col2 = st.columns(2)
        default_index = 0
        if selected_proj.logistic_status in status_opts: default_index = status_opts.index(selected_proj.logistic_status)
        elif selected_proj.logistic_status: default_index = 3
            
        selected_status = col1.selectbox("Status Logistik", status_opts, index=default_index)
        custom_val = selected_proj.logistic_status if default_index == 3 else ""
        custom_status = col2.text_input("Status Logistik (Custom)", value=custom_val) if selected_status == "Lainnya (Ketik Bebas)" else ""
        
        delivery_number = st.text_input("Nomor Surat Jalan (Delivery Number)", value=selected_proj.delivery_number or "")
        actual_date = st.date_input("Tanggal Aktual Pengiriman / Diterima", value=selected_proj.actual_delivery_date or get_jakarta_time().date())
        
        if st.form_submit_button("Simpan Data Logistik"):
            final_status = custom_status if selected_status == "Lainnya (Ketik Bebas)" else selected_status
            selected_proj.logistic_status = final_status
            selected_proj.delivery_number = delivery_number
            selected_proj.actual_delivery_date = actual_date
            db.commit()
            st.success(f"✅ Status Logistik & Surat Jalan untuk SO {selected_proj.so_number} berhasil diperbarui!")
            time.sleep(1.5); st.rerun()
    db.close()

# --- FUNGSI GENERATE EXCEL (MODULE 5) ---
def generate_excel_resume(project, logs):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Resume {project.so_number.replace('#', '')}"

    # 1. HEADER INFO
    ws["A1"] = "PROJECT RESUME - SMART MES & TONNAGE TRACKER"
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:G1")

    ws["A3"] = "SO Number"; ws["B3"] = project.so_number
    ws["A4"] = "Nama Proyek"; ws["B4"] = project.project_name
    ws["A5"] = "Total Tonase Target"; ws["B5"] = f"{project.total_tonnage} Kg/Ton"
    ws["A6"] = "Tanggal Mulai Produksi"; ws["B6"] = project.start_date.strftime("%d/%m/%Y") if project.start_date else "-"
    ws["A7"] = "Target Selesai Produksi"; ws["B7"] = project.target_selesai_global.strftime("%d/%m/%Y") if project.target_selesai_global else "-"
    ws["A8"] = "Status Logistik"; ws["B8"] = project.logistic_status if project.logistic_status else "-"
    ws["A9"] = "Nomor Surat Jalan"; ws["B9"] = project.delivery_number if project.delivery_number else "-"
    ws["A10"] = "Siklus YEJ"; ws["B10"] = project.yej_cycle if project.yej_cycle else "-"
    ws["A11"] = "Jenis Material"; ws["B11"] = project.material_list if project.material_list else "-"
    ws["A12"] = "Jenis Kontrak"; ws["B12"] = "🛠️ JASA ONLY" if project.is_jasa_only else "📦 MATERIAL + JASA"
    ws["A13"] = "Analisis Performa Waktu"; ws["B13"] = project.time_performance_status

    # 2. TABEL DURASI TAHAPAN
    start_row_dur = 15
    ws[f"A{start_row_dur}"] = "TRACKING DURASI WAKTU PER TAHAPAN (SCOPE)"
    ws[f"A{start_row_dur}"].font = Font(bold=True)
    
    headers_dur = ["Tahapan (Scope)", "Start Date", "End Date", "Total Durasi (Days)"]
    ws.append(headers_dur)
    h_row_dur = ws.max_row
    for col_num in range(1, 5):
        cell = ws.cell(row=h_row_dur, column=col_num)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="107C41", fill_type="solid")

    db = SessionLocal()
    active_routings = db.query(ProjectScopeRouting).filter(ProjectScopeRouting.project_id == project.id).all()
    for r in active_routings:
        sc_logs = [l for l in logs if l.scope_id == r.scope_id]
        if sc_logs:
            start_dt = min(l.log_date for l in sc_logs)
            pass_logs = [l for l in sc_logs if l.status == QCStatus.PASS]
            end_dt = max(l.log_date for l in pass_logs) if pass_logs else max(l.log_date for l in sc_logs)
            dur_days = max((end_dt.date() - start_dt.date()).days, 1)
            ws.append([r.scope.name, start_dt.strftime("%d/%m/%Y"), end_dt.strftime("%d/%m/%Y"), f"{dur_days} Hari"])
        else:
            ws.append([r.scope.name, "-", "-", "-"])
    db.close()

    ws.append([]) 

    # 3. TABEL DETAIL LOG LAPANGAN
    ws.append(["DETAIL LOG PROGRESS LAPANGAN"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)

    headers = ["Tanggal Log", "Tahapan (Scope)", "Vendor Pelaksana", "Tonase Aktual", "Status QC", "Catatan Temuan QC", "Foto Bukti", "Keterangan Foto"]
    ws.append(headers) 
    
    header_row = ws.max_row
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="003366", fill_type="solid")
        
    for log in logs:
        vendor_text = log.vendor.name if log.vendor else "Unknown Vendor"
        if getattr(log, 'po_vendor_number', None):
            vendor_text += f" (PO: {log.po_vendor_number})"

        caption = getattr(log, 'image_caption', None)
        final_caption = caption if caption else "-"

        row_data = [
            log.log_date.strftime("%d/%m/%Y %H:%M"), 
            log.scope.name, 
            vendor_text, 
            log.actual_tonnage_completed, 
            log.status.value, 
            log.qc_notes,
            "", 
            final_caption 
        ]
        ws.append(row_data)
        current_row = ws.max_row
        
        if log.image_path and os.path.exists(log.image_path):
            try:
                img = ExcelImage(log.image_path)
                img.height = 100; img.width = 100 
                ws.add_image(img, f"G{current_row}")
                ws.row_dimensions[current_row].height = 80
            except:
                ws[f"G{current_row}"] = "Error Image"
        else:
            ws[f"G{current_row}"] = "-"
            
    for col in ws.columns:
        max_length = 0
        column = get_column_letter(col[0].column)
        if column == 'G': 
            ws.column_dimensions[column].width = 15; continue
        for cell in col:
            try: max_length = max(max_length, len(str(cell.value)))
            except: pass
        ws.column_dimensions[column].width = (max_length + 2)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# --- MODULE 5: MANAGEMENT / PPIC (Dashboard Eksekutif) ---
def view_dashboard():
    st.header("📊 Management Dashboard & S-Curve Tonnage Base")
    db = SessionLocal()
    
    all_projects = db.query(Project).all()
    if not all_projects:
        st.info("Belum ada proyek yang terdaftar di sistem."); db.close(); return

    # FITUR 2: KARTU METRIK DASHBOARD
    total_so = len(all_projects)
    total_finished = sum(1 for p in all_projects if p.status_proyek == ProjectStatus.COMPLETED)
    total_delivered = sum(1 for p in all_projects if p.delivery_number and p.logistic_status and any(k in p.logistic_status.upper() for k in ['DIKIRIM', 'DITERIMA']))
    total_progress = sum(1 for p in all_projects if p.status_proyek in [ProjectStatus.IN_PRODUCTION, ProjectStatus.WAITING_DRAWING])

    c1, c2, c3, c4 = st.columns(4)
    c1.markdown(f"<div class='metric-card'><h4>Total SO</h4><h2>{total_so}</h2></div>", unsafe_allow_html=True)
    c2.markdown(f"<div class='metric-card'><h4>Total Progress</h4><h2>{total_progress}</h2></div>", unsafe_allow_html=True)
    c3.markdown(f"<div class='metric-card'><h4>Total Finished</h4><h2>{total_finished}</h2></div>", unsafe_allow_html=True)
    c4.markdown(f"<div class='metric-card'><h4>Total Delivered</h4><h2>{total_delivered}</h2></div>", unsafe_allow_html=True)

    st.subheader("📋 Ringkasan Progres Proyek Global")
    summary_data = []
    for proj in all_projects:
        logs = db.query(ProjectProgressLog).filter(ProjectProgressLog.project_id == proj.id).all()
        routings_count = db.query(ProjectScopeRouting).filter(ProjectScopeRouting.project_id == proj.id).count()
        total_target_work = proj.total_tonnage * routings_count if routings_count > 0 else 0
        
        total_actual_done = sum(l.actual_tonnage_completed for l in logs)
        avg_progress = min(max((total_actual_done / total_target_work) if total_target_work > 0 else 0.0, 0.0), 1.0)
        
        summary_data.append({
            "SO Number": proj.so_number, 
            "Nama Proyek": proj.project_name, 
            "Target Tonase": proj.total_tonnage,
            "Rata-Rata Progres": avg_progress,
            "Status Logistik": proj.logistic_status,
            "Performa Waktu": proj.time_performance_status, 
            "Status": proj.status_proyek.value
        })

    st.dataframe(pd.DataFrame(summary_data), column_config={
            "Rata-Rata Progres": st.column_config.ProgressColumn("Rata-Rata Progres", format="%.2f", min_value=0.0, max_value=1.0),
        }, use_container_width=True, hide_index=True)

    st.markdown("---")
    st.subheader("📈 Analisis Kurva S, Tracking Durasi & Download Excel")
    active_projects = [p for p in all_projects if p.status_proyek in [ProjectStatus.IN_PRODUCTION, ProjectStatus.COMPLETED]]
    
    if active_projects:
        colA, colB = st.columns([3, 1])
        selected_scurve_proj = colA.selectbox("Pilih Proyek untuk Kurva S & Laporan", active_projects, format_func=lambda x: f"{x.so_number} - {x.project_name}")
        
        proj_logs = db.query(ProjectProgressLog).filter(ProjectProgressLog.project_id == selected_scurve_proj.id, ProjectProgressLog.status == QCStatus.PASS).order_by(ProjectProgressLog.log_date.asc()).all()
        all_proj_logs = db.query(ProjectProgressLog).filter(ProjectProgressLog.project_id == selected_scurve_proj.id).order_by(ProjectProgressLog.log_date.asc()).all()
        
        with colB:
            st.write(""); st.write("")
            excel_file = generate_excel_resume(selected_scurve_proj, all_proj_logs)
            st.download_button(label="📥 Download Laporan Excel", data=excel_file, file_name=f"Resume_Proyek_{selected_scurve_proj.so_number}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # FITUR 3: TRACKING DURASI DISPLAY ON UI
        st.write("**⏱️ Detail Durasi Proses Per Tahapan (Mulai - Selesai):**")
        scope_durations = []
        active_routings = db.query(ProjectScopeRouting).filter(ProjectScopeRouting.project_id == selected_scurve_proj.id).all()
        for r in active_routings:
            sc_logs = [l for l in all_proj_logs if l.scope_id == r.scope_id]
            if sc_logs:
                start_dt = min(l.log_date for l in sc_logs)
                pass_logs = [l for l in sc_logs if l.status == QCStatus.PASS]
                end_dt = max(l.log_date for l in pass_logs) if pass_logs else max(l.log_date for l in sc_logs)
                dur_days = max((end_dt.date() - start_dt.date()).days, 1)
                scope_durations.append({
                    "Tahapan (Scope)": r.scope.name, "Mulai (Start Date)": start_dt.strftime("%d/%m/%Y"),
                    "Selesai (End Date)": end_dt.strftime("%d/%m/%Y"), "Total Durasi": f"{dur_days} Hari"
                })
            else: scope_durations.append({"Tahapan (Scope)": r.scope.name, "Mulai (Start Date)": "-", "Selesai (End Date)": "-", "Total Durasi": "-"})
        st.dataframe(pd.DataFrame(scope_durations), use_container_width=True, hide_index=True)

        total_target_work = selected_scurve_proj.total_tonnage
        start_date = selected_scurve_proj.start_date
        end_date = selected_scurve_proj.target_selesai_global
        
        if start_date and end_date and end_date >= start_date:
            date_range = pd.date_range(start=start_date, end=end_date)
            df_chart = pd.DataFrame(index=date_range)
            df_chart.index.name = 'date'
            n_days = len(date_range)
            s_curve_baseline = [total_target_work / (1 + math.exp(-(-6 + (12 * i / max(1, n_days - 1))))) for i in range(n_days)]
            df_chart['Target Estimasi Proyek'] = s_curve_baseline
            
            if proj_logs:
                actual_df = pd.DataFrame([{"date": pd.to_datetime(l.log_date.date()), "tonnage": l.actual_tonnage_completed} for l in proj_logs])
                actual_df = actual_df.groupby('date').sum()
                min_date = min(actual_df.index.min(), df_chart.index.min())
                full_date_range = pd.date_range(start=min_date, end=end_date)
                full_actual_df = pd.DataFrame(index=full_date_range)
                full_actual_df = full_actual_df.join(actual_df, how='left')
                full_actual_df['tonnage'] = full_actual_df['tonnage'].fillna(0).cumsum()
                df_chart = df_chart.join(full_actual_df['tonnage'], how='left')
                df_chart.rename(columns={'tonnage': 'Akumulasi Aktual'}, inplace=True)
                df_chart['Akumulasi Aktual'] = df_chart['Akumulasi Aktual'].ffill().fillna(0)
            else: df_chart['Akumulasi Aktual'] = 0.0
                
            st.line_chart(df_chart[['Target Estimasi Proyek', 'Akumulasi Aktual']], height=450)
        else:
            st.info("⚠️ Rentang waktu belum valid. PPIC harus menetapkan 'Tanggal Mulai' dan 'Target Selesai'.")
    else: st.info("Belum ada proyek berstatus 'In Production' atau 'Completed'.")
    db.close()

# --- STREAMLIT ROUTER MAIN ---
if "logged_in" not in st.session_state: st.session_state["logged_in"] = False

if not st.session_state["logged_in"]:
    login_screen()
else:
    with st.sidebar:
        display_logo()
        st.title("Navigasi Sistem")
        st.info(f"👤 Hak Akses Aktif:\n**{st.session_state['role']}**")
        if st.button("Logout", type="primary", use_container_width=True): logout()
        st.markdown("---")
        
        allowed_menus = st.session_state["access"]
        # FITUR 1: Management Dashboard default view
        default_idx = allowed_menus.index("Management Dashboard") if "Management Dashboard" in allowed_menus else 0
        menu = st.radio("Pilih Modul:", allowed_menus, index=default_idx)
    
    if menu == "Registrasi Proyek": view_registrasi_proyek()
    elif menu == "Perencanaan Target": view_perencanaan_target()
    elif menu == "Gambar Kerja": view_gambar_kerja()
    elif menu == "Update Progress & QC": view_update_qc()
    elif menu == "Update Status Logistik": view_update_logistik()
    elif menu == "Management Dashboard": view_dashboard()
